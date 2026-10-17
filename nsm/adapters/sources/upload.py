"""직접 올린 자료(가장 믿을 만한 1차 자료). 폴더의 파일을 읽어 자료 항목으로 만든다.

파일 경로 + 크기 + 수정 시각으로 해시를 만들어, 파일을 고치면 새 자료로 다시 들어간다.
이미지(png·jpg·webp)는 이미지 읽기(Claude)로 보이는 글자·사실을 옮겨 자료로 쓴다(C16, ADR 0045).
"""

from __future__ import annotations

import csv
import logging
from collections.abc import Callable
from pathlib import Path

from nsm.adapters.sources.common import clip, url_hash
from nsm.adapters.sources.image_reader import CACHE_DIR_NAME, IMAGE_SUFFIXES, is_image
from nsm.domain.clock import new_id
from nsm.domain.models import SourceItem
from nsm.ports import Clock

log = logging.getLogger(__name__)

TEXT_SUFFIXES = {".txt", ".md", ".csv", ".xlsx", ".pdf"}
SUPPORTED = TEXT_SUFFIXES | IMAGE_SUFFIXES

ImageReader = Callable[[Path], str]


class NoImageReader(ValueError):
    """이미지를 읽을 수단(Claude API 키)이 없음."""


def extract_text(path: Path, read_image: ImageReader | None = None) -> str:
    suffix = path.suffix.lower()
    if suffix in IMAGE_SUFFIXES:
        if read_image is None:
            raise NoImageReader("이미지를 읽으려면 Claude API 키가 필요합니다(`nsm init`)")
        return read_image(path)
    if suffix in (".txt", ".md"):
        return path.read_text(encoding="utf-8", errors="replace")
    if suffix == ".csv":
        with path.open(encoding="utf-8-sig", errors="replace", newline="") as f:
            return "\n".join(", ".join(row) for row in csv.reader(f))
    if suffix == ".xlsx":
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        lines = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                cells = [str(c) for c in row if c is not None]
                if cells:
                    lines.append(", ".join(cells))
        wb.close()
        return "\n".join(lines)
    if suffix == ".pdf":
        from pypdf import PdfReader

        return "\n".join((p.extract_text() or "") for p in PdfReader(path).pages)
    raise ValueError(f"지원하지 않는 파일 형식: {path.name}")


class UploadSource:
    name = "upload"

    def __init__(
        self,
        folder: Path,
        clock: Clock,
        is_known: Callable[[str], bool] | None = None,
        read_image: ImageReader | None = None,
    ) -> None:
        self.folder = folder
        self.clock = clock
        # 이미 저장된 파일(해시)은 건너뛰어야 상한에 막혀 뒤쪽 파일이 영영 안 들어오는 일이 없다
        self.is_known = is_known or (lambda _h: False)
        self.read_image = read_image

    def collect(self, topics: list[str], limit: int) -> list[SourceItem]:
        if not self.folder.is_dir():
            return []
        items: list[SourceItem] = []
        for path in sorted(self.folder.rglob("*")):
            if len(items) >= limit:
                break
            if not path.is_file() or path.suffix.lower() not in SUPPORTED:
                continue
            if CACHE_DIR_NAME in path.relative_to(self.folder).parts:
                continue  # 이미지에서 옮긴 글 캐시는 자료가 아니다
            st = path.stat()
            uri = path.resolve().as_uri()
            h = url_hash(f"{uri}#{st.st_size}-{int(st.st_mtime)}")
            if self.is_known(h) or (is_image(path) and self.read_image is None):
                continue  # 키가 없으면 이미지는 조용히 건너뛴다(올릴 때 화면에서 안내)
            try:
                text = clip(extract_text(path, self.read_image))
            except Exception as e:  # noqa: BLE001  파일 하나가 깨져도 나머지는 계속
                log.warning("자료 파일을 읽지 못했습니다 %s: %s", path, e)
                continue
            if not text.strip():
                continue
            items.append(
                SourceItem(
                    id=new_id(),
                    source=self.name,
                    url=uri,
                    title=path.stem,
                    text=text,
                    collected_at=self.clock.now(),
                    url_hash=h,
                )
            )
        return items
